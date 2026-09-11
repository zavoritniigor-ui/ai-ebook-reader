"""Optional fixture regeneration; CI reads the committed files, no dependencies.

Run with python-docx==1.2.0 and Pillow==12.3.0 installed. These are original,
redistributable test documents, with fixed timestamps and no private content.
"""
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
import zipfile

from docx import Document
from docx.shared import Inches
from PIL import Image, ImageDraw

DEST = Path(__file__).parent / 'fixtures' / 'text-audit'
TEXT = 'Hello reader. The little bird sings clearly. Привіт, світе! Zażółć gęślą jaźń. Café café 👩🏽‍💻 — it’s readable.'


def picture(kind):
    img = Image.new('RGB', (320, 220), '#1570b8')
    draw = ImageDraw.Draw(img)
    draw.rectangle((25, 25, 145, 195), fill='#e5b638')
    draw.ellipse((175, 45, 285, 155), fill='#d03842')
    buf = BytesIO()
    img.save(buf, format=kind)
    return buf.getvalue()


def save(doc, name):
    doc.core_properties.created = datetime(2020, 1, 1, tzinfo=timezone.utc)
    doc.core_properties.modified = datetime(2020, 1, 1, tzinfo=timezone.utc)
    doc.core_properties.author = 'AI Ebook Reader deterministic audit'
    out = BytesIO()
    doc.save(out)
    with zipfile.ZipFile(out) as source, zipfile.ZipFile(DEST / name, 'w', zipfile.ZIP_DEFLATED) as dest:
        for path in sorted(source.namelist()):
            info = zipfile.ZipInfo(path, date_time=(2020, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            dest.writestr(info, source.read(path))


def main():
    DEST.mkdir(exist_ok=True)
    for kind in ('PNG', 'JPEG'):
        (DEST / ('pattern.' + ('png' if kind == 'PNG' else 'jpg'))).write_bytes(picture(kind))
    for variant in ('text', 'mixed', 'images', 'almost-empty', 'long', 'empty'):
        doc = Document()
        if variant in ('text', 'mixed', 'long'):
            doc.add_heading('Audit chapter one', level=1)
            doc.add_paragraph(TEXT)
            p = doc.add_paragraph('A paragraph with ')
            p.add_run('bold text').bold = True
            p.add_run(' and ')
            p.add_run('italic text').italic = True
            doc.add_paragraph('Bullet one', style='List Bullet')
            doc.add_paragraph('Bullet two', style='List Bullet')
            doc.add_paragraph('Number one', style='List Number')
            doc.add_paragraph('Number two', style='List Number')
            table = doc.add_table(rows=2, cols=2)
            for cell, text in zip([cell for row in table.rows for cell in row.cells], ['Language', 'Text', 'Polski', 'Zażółć']):
                cell.text = text
        if variant in ('mixed', 'images'):
            for n in range(4):
                kind = 'JPEG' if n % 2 else 'PNG'
                doc.add_picture(BytesIO(picture(kind)), width=Inches(3))
                if variant == 'mixed':
                    doc.add_paragraph('Illustration caption ' + str(n + 1), style='Caption')
        if variant in ('text', 'mixed', 'long'):
            doc.add_page_break()
            doc.add_heading('Audit chapter two', level=1)
            doc.add_paragraph('Second chapter. ' + TEXT)
        if variant == 'long':
            for n in range(220):
                doc.add_paragraph(f'Paragraph {n}. ' + TEXT * 5)
        if variant == 'almost-empty':
            doc.add_paragraph('A')
        save(doc, variant + '.docx')


if __name__ == '__main__':
    main()
